# -*- coding: utf-8 -*-
"""路特创新专属 · AI算法测试工程师面试题库 · 400题"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

DATA = []
def S(title):
    DATA.append(("__SECTION__", title, ""))
def Q(q, a):
    DATA.append(("__QA__", q, a))

print("正在加载 400 题数据...")
# 数据由 gen_lute_data.py 生成
exec(open(os.path.join(os.path.dirname(__file__), 'gen_lute_data.py'), encoding='utf-8').read())

print(f"全部 {sum(1 for x in DATA if x[0]=='__QA__')} 题加载完成！")

# 生成 Word
print("\n正在生成 Word 文档...")
doc = Document()
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# 封面
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('路特创新 · AI算法测试工程师面试题库')
r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor(0,51,102)
r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('400题专属定制版\n\n岗位：AI算法测试工程师（35-52K·15薪）\n候选人：黄继根\n生成日期：2025年6月26日')
r2.font.size = Pt(12); r2.font.color.rgb = RGBColor(100,100,100)
r2.font.name = '微软雅黑'; r2._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

doc.add_page_break()

# 目录
toc = doc.add_paragraph()
toc_run = toc.add_run('目  录')
toc_run.bold = True; toc_run.font.size = Pt(18); toc_run.font.color.rgb = RGBColor(0,51,102)
toc_run.font.name = '微软雅黑'; toc_run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

chapters = [
    ("第一章", "大模型基础原理（Transformer/训练/微调）", "Q1-Q30"),
    ("第二章", "Prompt工程与评测方法", "Q31-Q80"),
    ("第三章", "RAG与检索增强", "Q81-Q120"),
    ("第四章", "Agent与工具调用", "Q121-Q170"),
    ("第五章", "模型性能与推理优化", "Q171-Q220"),
    ("第六章", "安全与对抗测试", "Q221-Q270"),
    ("第七章", "数据与评测体系", "Q271-Q310"),
    ("第八章", "AI代码生成测试", "Q311-Q340"),
    ("第九章", "路特创新专属场景题", "Q341-Q370"),
    ("第十章", "综合能力与项目实战", "Q371-Q400"),
]
for ch, title, qr in chapters:
    tp = doc.add_paragraph()
    tr = tp.add_run(f'{ch}  {title}')
    tr.bold = True; tr.font.size = Pt(12); tr.font.name = '微软雅黑'
    tr._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    tr2 = tp.add_run(f'\n    {qr}')
    tr2.font.size = Pt(10); tr2.font.color.rgb = RGBColor(120,120,120)
    tr2.font.name = '微软雅黑'; tr2._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

doc.add_page_break()

# 题目内容
qa_count = 0
for item in DATA:
    qtype, q, a = item
    if qtype == "__SECTION__":
        if q.startswith("第"):
            doc.add_page_break()
        p = doc.add_paragraph()
        r = p.add_run(q)
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0,51,102)
        r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    elif qtype == "__QA__":
        qa_count += 1
        pq = doc.add_paragraph()
        rq = pq.add_run(q)
        rq.bold = True; rq.font.size = Pt(11); rq.font.color.rgb = RGBColor(30,30,30)
        rq.font.name = '微软雅黑'; rq._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        pa = doc.add_paragraph()
        ra = pa.add_run(a)
        ra.font.size = Pt(10.5); ra.font.color.rgb = RGBColor(60,60,60)
        ra.font.name = '微软雅黑'; ra._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        ps = doc.add_paragraph()
        rs = ps.add_run('─' * 60)
        rs.font.size = Pt(8); rs.font.color.rgb = RGBColor(200,200,200)
        rs.font.name = '微软雅黑'; rs._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

OUTDIR = os.path.join(os.path.dirname(__file__), '..', 'docs')
os.makedirs(OUTDIR, exist_ok=True)
DOCX_FILE = os.path.join(OUTDIR, '路特创新-AI算法测试工程师面试题库-400题.docx')
doc.save(DOCX_FILE)
print(f"[OK] Word: {DOCX_FILE} | {qa_count}题 | {os.path.getsize(DOCX_FILE)/1024/1024:.2f}MB")

# 生成 PDF
print("\n正在生成 PDF...")
from fpdf import FPDF

class PDF(FPDF):
    def __init__(self):
        super().__init__('P', 'mm', 'A4')
        font_dir = os.path.join(os.environ.get('WINDIR', 'C:\\Windows'), 'Fonts')
        msyh = os.path.join(font_dir, 'msyh.ttc')
        msyhbd = os.path.join(font_dir, 'msyhbd.ttc')
        if os.path.exists(msyh):
            self.add_font('MSYH', '', msyh)
        else:
            self.add_font('MSYH', '', os.path.join(font_dir, 'simhei.ttf'))
        if os.path.exists(msyhbd):
            self.add_font('MSYH', 'B', msyhbd)
        else:
            self.add_font('MSYH', 'B', msyh)
        self.set_auto_page_break(True, 20)

    def header(self):
        if self.page_no() > 2:
            self.set_font('MSYH', '', 8)
            self.set_text_color(150,150,150)
            self.cell(0,5,'路特创新 · AI算法测试工程师面试题库 - 400题', align='L')
            self.cell(0,5,f'第{self.page_no()}页', align='R', new_x="LMARGIN", new_y="NEXT")
            self.line(self.l_margin, self.get_y(), self.w-self.r_margin, self.get_y())
            self.ln(3)

pdf = PDF()
pdf.set_title('路特创新-AI算法测试工程师面试题库-400题')
pdf.set_author('黄继根')

# 封面
pdf.add_page()
pdf.ln(45)
pdf.set_font('MSYH', 'B', 26)
pdf.set_text_color(0,51,102)
pdf.cell(0,14,'路特创新', align='C', new_x="LMARGIN", new_y="NEXT")
pdf.set_font('MSYH', 'B', 22)
pdf.cell(0,12,'AI算法测试工程师面试题库', align='C', new_x="LMARGIN", new_y="NEXT")
pdf.ln(6)
pdf.set_font('MSYH', '', 14)
pdf.set_text_color(100,100,100)
pdf.cell(0,10,'400题专属定制版', align='C', new_x="LMARGIN", new_y="NEXT")
pdf.ln(10)
pdf.set_font('MSYH', '', 11)
pdf.set_text_color(80,80,80)
pdf.cell(0,8,'岗位：AI算法测试工程师（35-52K·15薪）', align='C', new_x="LMARGIN", new_y="NEXT")
pdf.cell(0,8,'候选人：黄继根', align='C', new_x="LMARGIN", new_y="NEXT")
pdf.cell(0,8,'生成日期：2025年6月26日', align='C', new_x="LMARGIN", new_y="NEXT")

# 目录
pdf.add_page()
pdf.set_font('MSYH', 'B', 18)
pdf.set_text_color(0,51,102)
pdf.cell(0,12,'目  录', align='L', new_x="LMARGIN", new_y="NEXT")
pdf.ln(6)
for ch, title, qr in chapters:
    pdf.set_font('MSYH', 'B', 12)
    pdf.set_text_color(30,30,30)
    pdf.cell(0,8,f'{ch}  {title}', new_x="LMARGIN", new_y="NEXT")
    pdf.set_font('MSYH', '', 10)
    pdf.set_text_color(120,120,120)
    pdf.cell(0,6,f'     {qr}', new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

# 题目
qa_pdf = 0
for item in DATA:
    qtype, q, a = item
    if qtype == "__SECTION__":
        pdf.add_page()
        pdf.set_font('MSYH', 'B', 14)
        pdf.set_text_color(0,51,102)
        pdf.multi_cell(0,8,q, align='L')
        pdf.ln(5)
    elif qtype == "__QA__":
        qa_pdf += 1
        pdf.set_font('MSYH', 'B', 11)
        pdf.set_text_color(30,30,30)
        pdf.multi_cell(0,6,q, align='L')
        pdf.ln(1)
        pdf.set_font('MSYH', '', 10)
        pdf.set_text_color(60,60,60)
        pdf.multi_cell(0,5.5,a, align='L')
        pdf.set_draw_color(220,220,220)
        pdf.set_line_width(0.2)
        pdf.line(pdf.l_margin+5, pdf.get_y()+2, pdf.w-pdf.r_margin-5, pdf.get_y()+2)
        pdf.ln(5)

PDF_FILE = os.path.join(OUTDIR, '路特创新-AI算法测试工程师面试题库-400题.pdf')
pdf.output(PDF_FILE)
print(f"[OK] PDF: {PDF_FILE} | {qa_pdf}题 | {os.path.getsize(PDF_FILE)/1024/1024:.2f}MB")
print("\n[OK] 路特创新400题专属题库生成完毕！")
